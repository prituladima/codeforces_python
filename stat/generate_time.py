#!/usr/bin/env python3
"""Minutes-per-attempt chart from "Competitive programming.xlsx" (first sheet).

Usage:  python generate_time.py <type> <letter> <gate_minutes>
        python generate_time.py CF_div2 C 45        ->  times_for_CF_div2_C.png
        python generate_time.py Leetcode Medium 25  ->  times_for_Leetcode_Medium.png

Columns used: A contest name (filled on the first row of a group), B date, C type (CF_div2, CF_div3,
AtCoder_Beg, Leetcode, ...), D problem letter, F status ("Accepted" or "DNS"), H hours, I minutes, J seconds.
Rating (G) is not used.  Every point is time spent: a filled dot when Accepted, a hollow square when DNS.
A gate pass is an Accepted attempt within the gate (45:59 counts as 45).  Nothing is fetched from the internet.
"""
import sys, os, datetime, math
from openpyxl import load_workbook

HANDLE = 'prituladima'
WINDOW, NEED = 8, 6      # gate: NEED of the last WINDOW attempts accepted within the gate (45:59 counts as 45)

def load_rows(path, typ, letter):
    ws = load_workbook(path, data_only=True).worksheets[0]
    rows, contest = [], ''
    for r in ws.iter_rows(values_only=True):
        if r[0]:
            contest = str(r[0])
        if r[2] != typ or not str(r[3] or '').startswith(letter):
            continue
        if not isinstance(r[1], datetime.datetime):
            continue
        minutes = (r[7] or 0) * 60 + (r[8] or 0) + (r[9] or 0) / 60
        if minutes <= 0:
            continue
        status = str(r[5] or '').strip()
        rows.append(dict(date=r[1].strftime('%Y-%m-%d'), contest=contest, problem=str(r[3]),
                         solved=status.lower() == 'accepted', minutes=minutes))
    return rows

def short(contest):
    c = (contest.replace('Educational Codeforces Round ', 'Edu ')
                .replace('Codeforces Round ', 'R')
                .replace('AtCoder Beginner Contest ', 'ABC '))
    c = c.split(' (')[0]
    return c

def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def build_svg(rows, typ, letter, gate):
    W, H = 1200, 660
    PX0, PX1, PY0, PY1 = 70, 1160, 112, 520
    n = len(rows)
    YMAX = max(max(x['minutes'] for x in rows), gate)      # the slowest attempt sits on the top line
    slot = (PX1 - PX0) / n
    X = lambda i: PX0 + (i + 0.5) * slot
    Y = lambda m: PY1 - m / YMAX * (PY1 - PY0)
    SANS = 'font-family="Liberation Sans, Arial, sans-serif"'
    passes = lambda x: x['solved'] and x['minutes'] < gate + 1   # Accepted within the gate; 45:59 counts as 45
    last = rows[-WINDOW:]
    passed = sum(1 for x in last if passes(x))
    gate_ok = len(last) >= WINDOW and passed >= NEED
    status, scol = ('GATE PASSED', '#2e7d32') if gate_ok else ('GATE IS NOT PASSED', '#b8860b')
    # least-squares trend over all attempts
    xs = list(range(n)); ys = [x['minutes'] for x in rows]
    mx, my = sum(xs) / n, sum(ys) / n
    var = sum((a - mx) ** 2 for a in xs)
    slope = sum((a - mx) * (b - my) for a, b in zip(xs, ys)) / var if var else 0.0
    trend = lambda i: my + slope * (i - mx)
    clamp = lambda v: min(max(v, 0.0), YMAX)
    s = [f'<svg width="{W}" height="{H}" viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg">',
         f'<rect width="{W}" height="{H}" fill="#fcfcfb"/>',
         f'<text x="70" y="36" {SANS} font-size="18" font-weight="bold" fill="#0b0b0b">{esc(typ)} {esc(letter)}: minutes per attempt</text>',
         f'<text x="70" y="56" {SANS} font-size="12" fill="#52514e">{HANDLE} · {n} attempts, {rows[0]["date"]} to {rows[-1]["date"]}</text>',
         f'<text x="{PX1}" y="40" {SANS} font-size="15" font-weight="bold" fill="{scol}" text-anchor="end">{status}</text>',
         f'<text x="{PX1}" y="56" {SANS} font-size="11" fill="#52514e" text-anchor="end">{passed} of last {len(last)} within {gate:g} min, {NEED} of {WINDOW} needed</text>',
         f'<circle cx="{PX0 + 6}" cy="84" r="5" fill="#2a78d6"/>',
         f'<text x="{PX0 + 16}" y="88" {SANS} font-size="12" fill="#52514e">time spent, accepted</text>',
         f'<rect x="{PX0 + 150}" y="78" width="11" height="11" fill="#fcfcfb" stroke="#2a78d6" stroke-width="2"/>',
         f'<text x="{PX0 + 167}" y="88" {SANS} font-size="12" fill="#52514e">DNS</text>',
         f'<line x1="{PX0 + 215}" x2="{PX0 + 237}" y1="84" y2="84" stroke="#898781" stroke-width="2"/>',
         f'<text x="{PX0 + 243}" y="88" {SANS} font-size="12" fill="#52514e">trend, {slope:+.1f} min per attempt</text>']
    ticks = [v for v in range(0, int(YMAX) + 1, 30) if v <= YMAX - 12] + [YMAX]
    for v in ticks:
        y = Y(v)
        col = '#c3c2b7' if v == 0 else '#e1e0d9'
        s.append(f'<line x1="{PX0}" x2="{PX1}" y1="{y:.1f}" y2="{y:.1f}" stroke="{col}" stroke-width="1"/>')
        s.append(f'<text x="{PX0 - 8}" y="{y + 4:.1f}" {SANS} font-size="11" fill="#52514e" text-anchor="end">{round(v)}</text>')
    yg = Y(gate)
    s.append(f'<line x1="{PX0}" x2="{PX1}" y1="{yg:.1f}" y2="{yg:.1f}" stroke="#52514e" stroke-width="1.5" stroke-dasharray="6 5"/>')
    s.append(f'<text x="{PX1}" y="{yg - 6:.1f}" {SANS} font-size="11" fill="#52514e" text-anchor="end">{gate:g}-minute gate</text>')
    s.append(f'<line x1="{X(0):.1f}" y1="{Y(clamp(trend(0))):.1f}" x2="{X(n - 1):.1f}" y2="{Y(clamp(trend(n - 1))):.1f}" stroke="#898781" stroke-width="2"/>')
    pts = ' '.join(f'{X(i):.1f},{Y(x["minutes"]):.1f}' for i, x in enumerate(rows))
    s.append(f'<polyline points="{pts}" fill="none" stroke="#2a78d6" stroke-width="2" stroke-linejoin="round" stroke-linecap="round"/>')
    maxc = max(6, int(slot / 6.3))           # characters that fit under one slot at 11px
    for i, x in enumerate(rows):
        cx, cy, m = X(i), Y(x['minutes']), round(x['minutes'])
        if x['solved']:
            s.append(f'<circle cx="{cx:.1f}" cy="{cy:.1f}" r="6" fill="#2a78d6" stroke="#fcfcfb" stroke-width="2"/>')
        else:
            s.append(f'<rect x="{cx - 6:.1f}" y="{cy - 6:.1f}" width="12" height="12" fill="#fcfcfb" stroke="#2a78d6" stroke-width="2.5"/>')
        ly = cy + 20 if x['minutes'] < gate else cy - 13   # labels under the gate line go below the point
        s.append(f'<text x="{cx:.1f}" y="{ly:.1f}" {SANS} font-size="13" font-weight="bold" fill="#0b0b0b" text-anchor="middle">{m}</text>')
        s.append(f'<text x="{cx:.1f}" y="{PY1 + 26}" {SANS} font-size="13" font-weight="bold" fill="#0b0b0b" text-anchor="middle">{i + 1}</text>')
        name = short(x['contest']) + (f' {x["problem"]}' if x['problem'] != letter else '')
        if len(name) > maxc:
            name = name[:maxc - 1] + '…'
        s.append(f'<text x="{cx:.1f}" y="{PY1 + 42}" {SANS} font-size="11" fill="#898781" text-anchor="middle">{esc(name)}</text>')
    i = 0
    while i < n:
        j = i
        while j + 1 < n and rows[j + 1]['date'] == rows[i]['date']:
            j += 1
        x0, x1, yb = X(i) - slot / 2 + 8, X(j) + slot / 2 - 8, PY1 + 72
        s.append(f'<polyline points="{x0:.1f},{yb} {x0:.1f},{yb + 6} {x1:.1f},{yb + 6} {x1:.1f},{yb}" fill="none" stroke="#c3c2b7" stroke-width="1"/>')
        s.append(f'<text x="{(x0 + x1) / 2:.1f}" y="{yb + 22}" {SANS} font-size="11" fill="#52514e" text-anchor="middle">{rows[i]["date"]}</text>')
        i = j + 1
    s.append('</svg>')
    return '\n'.join(s)

def main():
    if len(sys.argv) != 4:
        sys.exit('usage: python generate_time.py <type> <letter> <gate_minutes>   e.g.  python generate_time.py CF_div2 C 45')
    typ, letter = sys.argv[1], sys.argv[2]
    try:
        gate = float(sys.argv[3])
    except ValueError:
        sys.exit(f'gate must be a number of minutes, got {sys.argv[3]!r}')
    if gate <= 0:
        sys.exit('gate must be positive')
    here = os.path.dirname(os.path.abspath(__file__))
    xlsx = os.path.join(here, 'Competitive programming.xlsx')
    rows = load_rows(xlsx, typ, letter)
    if not rows:
        sys.exit(f'no rows of type {typ} with problem {letter} in {xlsx}')
    svg = build_svg(rows, typ, letter, gate)
    base = os.path.join(here, f'times_for_{typ}_{letter}')
    try:
        import cairosvg
    except ImportError:
        with open(base + '.svg', 'w', encoding='utf-8') as f:
            f.write(svg)
        sys.exit(f'cairosvg not installed (conda env cf-stat): wrote {base}.svg only')
    cairosvg.svg2png(bytestring=svg.encode('utf-8'), write_to=base + '.png', output_width=2400)
    print(f'{len(rows)} attempts -> {base}.png')

if __name__ == '__main__':
    main()
